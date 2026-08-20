import os
import docx
from llama_index.core import Document
from llama_index.core.node_parser import SentenceSplitter

def process_docx_document(file_path: str, doc_name: str, ingestion_id: str, knowledge_id: str | None = None):
    print(f"Extracting DOCX with structure preservation: {doc_name}")
    knowledge_id = knowledge_id or doc_name
    doc = docx.Document(file_path)
    
    full_text = []
    
    for element in doc.element.body:
        if element.tag.endswith('p'):
            for p in doc.paragraphs:
                if p._element == element:
                    if p.style.name.startswith('Heading'):
                        full_text.append(f"\\n# {p.text}\\n")
                    else:
                        full_text.append(p.text)
                    break
        elif element.tag.endswith('tbl'):
            for table in doc.tables:
                if table._element == element:
                    table_text = []
                    for row in table.rows:
                        row_text = " | ".join(cell.text.replace("\\n", " ").strip() for cell in row.cells)
                        table_text.append(row_text)
                    full_text.append("\\n[TABLE]\\n" + "\\n".join(table_text) + "\\n[/TABLE]\\n")
                    break
                    
    combined_text = "\\n".join(full_text)
    
    nodes = []
    document = Document(text=combined_text)
    
    text_splitter = SentenceSplitter(chunk_size=512, chunk_overlap=50)
    doc_nodes = text_splitter.get_nodes_from_documents([document])
    
    for node in doc_nodes:
        node.metadata["repo_name"] = knowledge_id
        node.metadata["ingestion_id"] = ingestion_id
        node.metadata["file_name"] = doc_name
        node.metadata["file_path"] = doc_name
        node.metadata["source_type"] = "document"
        node.metadata["extension"] = ".docx"
        
        nodes.append(node)
        
    print(f"Successfully chunked {doc_name} into {len(nodes)} nodes.")
    return nodes
